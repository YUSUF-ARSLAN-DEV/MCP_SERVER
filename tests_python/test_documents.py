import json
import logging
from types import SimpleNamespace

import pytest

from website_test_pipeline.documents import (
    CHUNK_CHARS, MAX_CHUNKS, DocumentError, chunks, prompt_for, quote_in_text, read_document, run_documents,
)
from website_test_pipeline.intents import accept_intents, load_intents

LOG = logging.getLogger("docs")
REQ = ("Visitors must be able to search for satellite frequencies by choosing a country. "
       "After a search the results page shows the frequencies for that country.\n\n"
       "Visitors can subscribe to frequency change updates from the subscribe page.")


# ------------------------------------------------------------------ reading documents

def test_plain_text_and_markdown_are_read_as_they_are(tmp_path):
    (tmp_path / "a.md").write_text("# Requirements" + chr(10) + REQ, encoding="utf-8")
    (tmp_path / "b.txt").write_bytes(REQ.replace(chr(10), chr(13) + chr(10)).encode("utf-8"))
    assert "search for satellite frequencies" in read_document(tmp_path / "a.md")
    assert chr(13) not in read_document(tmp_path / "b.txt")                          # line endings are normalised


def test_a_docx_is_read_including_its_tables(tmp_path):
    from docx import Document
    document = Document()
    document.add_paragraph("Visitors can pick a country and search for the frequencies.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "REQ-7"
    table.rows[0].cells[1].text = "The subscribe form accepts a name and an email address."
    document.save(str(tmp_path / "spec.docx"))
    text = read_document(tmp_path / "spec.docx")
    assert "pick a country" in text and "REQ-7 | The subscribe form accepts a name" in text


def test_a_pdf_is_read(tmp_path):
    body = b"BT /F1 12 Tf 72 720 Td (Visitors can search for frequencies by country) Tj ET"
    pdf = (b"%PDF-1.4" + chr(10).encode() + b"1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj" + chr(10).encode()
           + b"2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj" + chr(10).encode()
           + b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>endobj" + chr(10).encode()
           + b"4 0 obj<</Length " + str(len(body)).encode() + b">>stream" + chr(10).encode() + body + chr(10).encode() + b"endstream endobj" + chr(10).encode()
           + b"5 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj" + chr(10).encode()
           + b"trailer<</Root 1 0 R/Size 6>>" + chr(10).encode() + b"startxref" + chr(10).encode() + b"0" + chr(10).encode() + b"%%EOF")
    (tmp_path / "spec.pdf").write_bytes(pdf)
    assert "search for frequencies by country" in read_document(tmp_path / "spec.pdf")


def test_documents_that_cannot_be_used_say_why(tmp_path):
    with pytest.raises(DocumentError, match="does not exist"):
        read_document(tmp_path / "nope.md")
    (tmp_path / "a.xlsx").write_bytes(b"x")
    with pytest.raises(DocumentError, match="not supported"):
        read_document(tmp_path / "a.xlsx")
    (tmp_path / "empty.txt").write_text("   " + chr(10) + chr(10), encoding="utf-8")
    with pytest.raises(DocumentError, match="no readable text"):
        read_document(tmp_path / "empty.txt")
    (tmp_path / "broken.docx").write_bytes(b"this is not a zip file")
    with pytest.raises(DocumentError, match="could not be read"):
        read_document(tmp_path / "broken.docx")


def test_a_huge_file_is_refused_before_it_is_read(tmp_path, monkeypatch):
    from website_test_pipeline import documents
    monkeypatch.setattr(documents, "MAX_BYTES", 10)
    (tmp_path / "big.txt").write_text("word " * 100, encoding="utf-8")
    with pytest.raises(DocumentError, match="larger than"):
        read_document(tmp_path / "big.txt")


# ------------------------------------------------------------------ chunking

def test_chunks_keep_paragraphs_whole_and_lose_nothing():
    paragraphs = [f"Paragraph {n} " + "word " * 40 for n in range(30)]
    text = (chr(10) * 2).join(paragraphs)
    parts = chunks(text, size=600)
    assert all(len(p) <= 600 for p in parts) and len(parts) > 1
    assert (chr(10) * 2).join(parts).replace(chr(10) * 2, " ").split() == text.replace(chr(10) * 2, " ").split()
    assert all(p.startswith("Paragraph") for p in parts)                          # cut between paragraphs, not inside them


def test_one_paragraph_longer_than_a_chunk_is_cut_and_a_short_text_is_one_chunk():
    assert chunks("a" * 250, size=100) == ["a" * 100, "a" * 100, "a" * 50]
    assert chunks("short text", size=100) == ["short text"] and chunks("", size=100) == []


# ------------------------------------------------------------------ the quote rule

def test_a_quote_must_really_be_in_the_document_but_punctuation_and_case_do_not_matter():
    assert quote_in_text("visitors must be able to search for satellite frequencies by choosing a country", REQ)
    assert quote_in_text("Visitors can subscribe to frequency change updates -- from the subscribe page!", REQ)
    assert quote_in_text("after a search, the RESULTS page shows the frequencies", REQ)


def test_a_paraphrase_an_invention_or_a_tiny_quote_is_never_accepted():
    assert not quote_in_text("Visitors are able to look up frequencies for their country", REQ)          # paraphrase
    assert not quote_in_text("The system must send a confirmation email after subscribing", REQ)         # invented requirement
    assert not quote_in_text("search for", REQ) and not quote_in_text("", REQ)                           # too short to prove anything
    assert not quote_in_text("search for satellite frequencies by country choosing", REQ)                # right words, wrong order


def test_the_quote_rule_works_for_any_script():
    arabic = "يجب أن يتمكن الزائر من البحث عن الترددات حسب الدولة"
    assert quote_in_text("يتمكن الزائر من البحث عن الترددات", arabic) and not quote_in_text("يتمكن الزائر من الاشتراك في القائمة", arabic)


# ------------------------------------------------------------------ accepting journeys from a document

ROW = {"sentence": "A visitor picks a country and searches, then sees the frequencies for that country.", "start_path": "/en",
       "evidence": "the home page has a country select", "quote": "the results page shows the frequencies for that country"}


def test_a_journey_with_a_real_quote_is_kept_with_its_provenance():
    doc = {"intents": []}
    added, rejected = accept_intents([ROW], doc, {"/en"}, "t", "m", document=("spec.md", REQ))
    assert not rejected and added[0]["source"] == "doc" and added[0]["from_doc"] == "spec.md"
    assert added[0]["quote"] == ROW["quote"] and added[0]["proposed_by"] == {"model": "m", "prompt_version": "intents-doc-v1"}


def test_an_invented_or_missing_quote_rejects_the_journey():
    doc = {"intents": []}
    invented = dict(ROW, quote="The site must support single sign-on for every visitor")
    missing = {k: v for k, v in ROW.items() if k != "quote"}
    added, rejected = accept_intents([invented, missing], doc, {"/en"}, "t", "m", document=("spec.md", REQ))
    assert added == [] and len(rejected) == 2 and all("is not in spec.md" in reason for _, reason in rejected)


def test_a_document_journey_is_held_to_the_same_rules_as_any_ai_sentence():
    doc = {"intents": []}
    account = dict(ROW, sentence="A visitor logs in with a password and then sees the frequencies page.")
    elsewhere = dict(ROW, start_path="/nowhere")
    added, rejected = accept_intents([account, elsewhere], doc, {"/en"}, "t", "m", document=("spec.md", REQ))
    reasons = " | ".join(r for _, r in rejected)
    assert added == [] and "credentials" in reasons and "not explored" in reasons


def test_without_a_document_the_quote_is_neither_required_nor_stored():
    added, _ = accept_intents([{k: v for k, v in ROW.items() if k != "quote"}], {"intents": []}, {"/en"}, "t", "m")
    assert added[0]["source"] == "ai" and "quote" not in added[0] and "from_doc" not in added[0]


# ------------------------------------------------------------------ the command

def _settings(tmp_path):
    artifacts = tmp_path / "artifacts"
    artifacts.mkdir()
    inv = {"url": "https://x.test/en", "title": "Home", "headings": [], "forms": [], "revealed": [], "embeds": [],
           "controls": [{"tag": "select", "name": "Country", "selector": "#c", "region": "other", "options": ["Egypt"]},
                        {"tag": "button", "name": "Search", "region": "other"}]}
    (artifacts / "x.inventory.json").write_text(json.dumps(inv), encoding="utf-8")
    return SimpleNamespace(intents_file=tmp_path / "i.json", flows_file=tmp_path / "f.json", artifacts_dir=artifacts,
                           urls_file=tmp_path / "u.txt", model="m")


class _Client:
    def __init__(self, *answers):
        self.answers, self.prompts = list(answers), []

    def generate(self, prompt, system):
        self.prompts.append(prompt)
        answer = self.answers.pop(0) if len(self.answers) > 1 else self.answers[0]
        if isinstance(answer, Exception):
            raise answer
        return answer if isinstance(answer, str) else json.dumps(answer)


URLS = ["https://x.test/en"]


def test_journeys_are_read_from_a_document_and_only_the_ones_resting_on_its_words_are_kept(tmp_path):
    settings = _settings(tmp_path)
    (tmp_path / "spec.md").write_text(REQ, encoding="utf-8")
    bad = dict(ROW, sentence="A visitor pays for a premium plan and then sees the dashboard page.", quote="Premium plans are paid monthly by card")
    client = _Client({"intents": [ROW, bad]})
    assert run_documents(settings, URLS, client, LOG, [str(tmp_path / "spec.md")]) == 0
    stored = load_intents(settings.intents_file)["intents"]
    assert [i["sentence"] for i in stored] == [ROW["sentence"]] and stored[0]["from_doc"] == "spec.md"
    assert "REQUIREMENTS (from spec.md, part 1 of 1)" in client.prompts[0] and "search for satellite frequencies" in client.prompts[0]
    assert "SITE MAP" in client.prompts[0] and "exact sentence or clause" in client.prompts[0]


def test_several_documents_are_read_in_turn_and_a_bad_one_does_not_stop_the_others(tmp_path):
    settings = _settings(tmp_path)
    (tmp_path / "spec.md").write_text(REQ, encoding="utf-8")
    client = _Client({"intents": [ROW]})
    code = run_documents(settings, URLS, client, LOG, [str(tmp_path / "missing.md"), str(tmp_path / "notes.xlsx"), str(tmp_path / "spec.md")])
    assert code == 0 and len(load_intents(settings.intents_file)["intents"]) == 1


def test_nothing_readable_or_no_useful_answer_is_reported_as_exit_1(tmp_path):
    settings = _settings(tmp_path)
    assert run_documents(settings, URLS, _Client({"intents": []}), LOG, [str(tmp_path / "missing.md")]) == 1
    (tmp_path / "spec.md").write_text(REQ, encoding="utf-8")
    assert run_documents(settings, URLS, _Client("no json at all"), LOG, [str(tmp_path / "spec.md")]) == 1
    assert not settings.intents_file.exists()


def test_no_explored_pages_is_exit_2_and_an_unavailable_model_keeps_what_was_added(tmp_path):
    empty = SimpleNamespace(intents_file=tmp_path / "i.json", flows_file=tmp_path / "f.json", artifacts_dir=tmp_path / "none",
                            urls_file=tmp_path / "u.txt", model="m")
    assert run_documents(empty, URLS, _Client({}), LOG, ["x.md"]) == 2

    from website_test_pipeline.llm import ModelError
    settings = _settings(tmp_path)
    big = (chr(10) * 2).join(f"Requirement {n}: " + "the visitor can do thing number %d on the site " % n * 30 for n in range(40))
    (tmp_path / "spec.md").write_text(REQ + chr(10) * 2 + big, encoding="utf-8")
    client = _Client({"intents": [ROW]}, ModelError("gateway timeout", status=524))
    assert run_documents(settings, URLS, client, LOG, [str(tmp_path / "spec.md")]) == 4
    assert len(load_intents(settings.intents_file)["intents"]) == 1                   # part 1's journey survived the outage in part 2


def test_a_long_document_is_read_in_a_bounded_number_of_parts(tmp_path):
    settings = _settings(tmp_path)
    long_text = (chr(10) * 2).join(f"Requirement {n}: " + "visitors must be able to do thing %d " % n * 40 for n in range(200))
    assert len(chunks(long_text)) > MAX_CHUNKS
    (tmp_path / "spec.md").write_text(long_text, encoding="utf-8")
    client = _Client({"intents": []})
    run_documents(settings, URLS, client, LOG, [str(tmp_path / "spec.md")])
    assert len(client.prompts) == MAX_CHUNKS and all(len(p) < CHUNK_CHARS + 20000 for p in client.prompts)


def test_a_sentence_already_known_is_not_added_twice_and_the_prompt_lists_it(tmp_path):
    settings = _settings(tmp_path)
    (tmp_path / "spec.md").write_text(REQ, encoding="utf-8")
    client = _Client({"intents": [ROW]})
    run_documents(settings, URLS, client, LOG, [str(tmp_path / "spec.md")])
    run_documents(settings, URLS, client, LOG, [str(tmp_path / "spec.md")])
    assert len(load_intents(settings.intents_file)["intents"]) == 1 and ROW["sentence"] in client.prompts[1]


def test_the_prompt_names_the_document_part_and_never_shows_more_than_a_chunk():
    text = prompt_for("SITE MAP", ["Existing sentence one"], "spec.docx", 2, 3, "excerpt text", "NOT YET COVERED by any tested flow")
    assert "REQUIREMENTS (from spec.docx, part 2 of 3)" in text and "excerpt text" in text
    assert "Existing sentence one" in text and "NOT YET COVERED by any tested flow" in text
    assert text.endswith("do not restate the requirements or explain first.")        # the last thing the model reads: JSON only
