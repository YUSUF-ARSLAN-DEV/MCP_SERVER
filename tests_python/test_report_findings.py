import json
from pathlib import Path

import pytest
from docx import Document
from PIL import Image

from website_test_pipeline.flowgen import file_name
from website_test_pipeline.report import _find_attachments, create_report, load_run

START = "https://x.test/en"


def _png(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (20, 20), "white").save(path)


def _text(path: Path) -> list[str]:
    doc = Document(str(path))
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        lines += [cell.text for row in table.rows for cell in row.cells]
    return lines


def _last_index(lines: list[str], text: str) -> int:
    """The real heading's position, not the Table of Contents entry of the same text near the top."""
    return len(lines) - 1 - lines[::-1].index(text)


def _flow(fid, status="verified", **over):
    flow = {"id": fid, "goal": "A visitor picks a country and sees the results.", "source": "intent", "status": status,
            "start_url": START, "steps": [{"kind": "click", "selector": None, "name": "Search"}],
            "outcome": {"effect": "navigates", "to": "/en/find"},
            "observed": {"effect": "navigates", "url": "https://x.test/en/find", "landed_url": START,
                         "step_urls": ["https://x.test/en/find"], "new_headings": ["Results"]}}
    flow.update(over)
    return flow


def _workspace(tmp_path):
    artifacts, tests = tmp_path / "artifacts", tmp_path / "tests"
    artifacts.mkdir()
    tests.mkdir()

    # a page-level spec whose failing assertion queries a role the explorer never recorded for that control
    page_spec = tests / "https-x-test-en_test.py"
    page_spec.write_text(
        "from playwright.sync_api import expect\n"
        "def test_wizard_step(page):\n"
        "    expect(page.get_by_role('group', name='Passcode (default)')).to_be_visible()\n", encoding="utf-8")
    (artifacts / "page.inventory.json").write_text(json.dumps(
        {"url": START, "controls": [{"name": "Passcode (default)", "tag": "fieldset", "role": None}]}), encoding="utf-8")

    good = _flow("f--good")
    flapper = _flow("f--flap", goal="A visitor searches and sees the flaky results page.")
    (tests / file_name(good)).write_text("def test_flow_good(page): pass\n", encoding="utf-8")
    (tests / file_name(flapper)).write_text("def test_flow_flap(page): pass\n", encoding="utf-8")

    page_node = "tests/https-x-test-en_test.py::test_wizard_step[chromium]"
    good_node = f"tests/{file_name(good)}::test_flow_good[chromium]"
    flap_node = f"tests/{file_name(flapper)}::test_flow_flap[chromium]"
    rows = [
        {"nodeid": page_node, "title": "wizard step", "url": START, "status": "failed", "duration": 1.0,
         "error": "E   AssertionError: Locator expected to be visible\nE     - waiting for get_by_role(\"group\", name=\"Passcode (default)\")"},
        {"nodeid": good_node, "title": "", "url": START, "status": "passed", "duration": 1.0, "error": None},
        {"nodeid": flap_node, "title": "", "url": START, "status": "passed", "duration": 1.0, "error": None},
    ]
    (artifacts / "test_results.json").write_text(json.dumps({"finished_at": "t", "tests": rows}), encoding="utf-8")
    _png(artifacts / "evidence" / "runs-tests-https-x-test-en-test-py-test-wizard-step-chromium" / "01.png")
    _png(artifacts / "evidence" / "runs-tests-flow-f-good-test-py-test-flow-good-chromium" / "01.png")
    _png(artifacts / "evidence" / "runs-tests-flow-f-flap-test-py-test-flow-flap-chromium" / "01.png")

    (tmp_path / "flows.json").write_text(json.dumps({"version": 1, "flows": [good, flapper]}), encoding="utf-8")
    (tmp_path / "flow_ratings.json").write_text(json.dumps({"version": 1, "ratings": {
        "f--good": [{"source": "runner", "at": "t1", "passed": True}, {"source": "runner", "at": "t2", "passed": True}],
        "f--flap": [{"source": "runner", "at": "t1", "passed": True},
                    {"source": "runner", "at": "t2", "passed": False, "error": "no content shown"},
                    {"source": "runner", "at": "t3", "passed": True}],
    }}), encoding="utf-8")
    return artifacts, tests


def test_load_run_populates_findings_and_flapping_from_real_files(tmp_path):
    artifacts, tests = _workspace(tmp_path)
    run = load_run(artifacts, tests)
    assert len(run.findings) == 1
    f = run.findings[0]
    assert f.severity == "Medium" and f.kind == "test_defect" and 'role="group"' in f.summary and "no role recorded" in f.summary
    assert (len(run.flapping) == 1 and run.flapping[0].test == "A visitor searches and sees the flaky results page."
            and run.flapping[0].sequence == "P F P")


def test_the_combined_report_leads_with_a_findings_section(tmp_path):
    artifacts, tests = _workspace(tmp_path)
    create_report(artifacts, tests, tmp_path / "report", combined=True)
    doc = Document(str(tmp_path / "report" / "full-report.docx"))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert headings.index("Defect Report (Bugs Found)") < headings.index("Test Coverage / Requirements Traceability")
    joined = chr(10).join(_text(tmp_path / "report" / "full-report.docx"))
    for part in ("Tooling issues", "Browser", "Operating system", "Python", "Low",
                 "1 test defect(s)", "Flow coverage", "Content controls a tested flow acts on",
                 "Tests run", "Journeys with no result", "Pass rate", "User journeys passed"):
        assert part in joined, part
    lines = _text(tmp_path / "report" / "full-report.docx")
    main = chr(10).join(lines[:_last_index(lines, "Test Logs & Evidence")])
    assert 'role="group"' not in main and "expect(" not in main


def test_reports_include_a_populated_clickable_table_of_contents(tmp_path):
    artifacts, tests = _workspace(tmp_path)
    create_report(artifacts, tests, tmp_path / "report", combined=True)
    doc = Document(str(tmp_path / "report" / "full-report.docx"))
    assert "Table of Contents" in [p.text for p in doc.paragraphs]
    assert "Executive Summary" in [p.text for p in doc.paragraphs]
    assert "Right-click this table" not in doc.part.element.xml
    assert doc.part.element.xml.count("w:hyperlink") >= 10


def test_the_consolidated_table_lists_every_test_with_expected_observed_and_verdict(tmp_path):
    artifacts, tests = _workspace(tmp_path)
    create_report(artifacts, tests, tmp_path / "report", combined=True)
    doc = Document(str(tmp_path / "report" / "full-report.docx"))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert headings.index("Test Execution Summary") < headings.index("Defect Report (Bugs Found)") < headings.index("Test Coverage / Requirements Traceability")
    table = [t for t in doc.tables if [c.text for c in t.rows[0].cells] == ["Test ID", "Scope", "Test name", "URL", "Expected", "Observed", "Result"]][-1]
    rows = [[c.text for c in r.cells] for r in table.rows[1:]]
    by_test = {r[2]: r for r in rows}
    assert by_test["wizard step"][1] == "page" and by_test["wizard step"][6] == "FAILED"
    assert "Passcode" in by_test["wizard step"][4]
    assert by_test["wizard step"][5] and by_test["wizard step"][5] != "As expected."
    flow_row = next(r for r in rows if r[2] == "A visitor picks a country and sees the results.")
    assert (flow_row[1] == "flow" and flow_row[4] == "The page navigates to a new address, landing on /en/find."
           and flow_row[6] == "PASSED")
    assert flow_row[5] and flow_row[5] != "(never ran)"


def test_a_wrong_test_is_a_test_defect_and_not_in_the_failure_count(tmp_path):
    artifacts, tests = _workspace(tmp_path)
    create_report(artifacts, tests, tmp_path / "report", combined=True)
    doc = Document(str(tmp_path / "report" / "full-report.docx"))
    # the fixture's failing page test asserts a role the page never had: the TEST is wrong, not the application
    assert not [t for t in doc.tables if [c.text for c in t.rows[0].cells][:2] == ["Test ID", "Defect"]]
    joined = chr(10).join(_text(tmp_path / "report" / "full-report.docx"))
    assert "0 application defects found this run." in joined
    assert "2 tests ran (0 page + 2 flow): 2 passed and 0 failed." in joined


def test_a_real_failure_is_a_defect_a_high_one_means_no_go_and_it_is_in_the_attention_table(tmp_path):
    artifacts, tests = _workspace(tmp_path)
    results = json.loads((artifacts / "test_results.json").read_text(encoding="utf-8"))
    for row in results["tests"]:
        if "test_flow_good" in row["nodeid"]:
            row.update(status="failed", error="E   AssertionError: Page URL expected to be 'https://x.test/en/find'\nE   Actual value: https://x.test/en")
    (artifacts / "test_results.json").write_text(json.dumps(results), encoding="utf-8")
    create_report(artifacts, tests, tmp_path / "report", combined=True)
    doc = Document(str(tmp_path / "report" / "full-report.docx"))
    joined = chr(10).join(_text(tmp_path / "report" / "full-report.docx"))
    assert joined.count("1 application defect found this run.") >= 1
    table = next(t for t in doc.tables if [c.text for c in t.rows[0].cells][:3] == ["ID", "Severity", "Priority"] and "Title" in [c.text for c in t.rows[0].cells])
    row = table.rows[1]
    assert row.cells[0].text == "DEF‑001" and row.cells[1].text == "High" and row.cells[2].text == "P1"
    assert "Release decision: NO-GO" in joined
    attention = next(t for t in doc.tables if [c.text for c in t.rows[0].cells][:2] == ["Test ID", "Defect"])
    assert [r.cells[7].text for r in attention.rows[1:]] == ["FAILED"]


def test_a_page_test_with_no_captured_assertion_says_so_plainly(tmp_path):
    artifacts, tests = _workspace(tmp_path)
    (tests / "https-x-test-en_test.py").write_text(
        "def test_wizard_step(page):\n    pass\n", encoding="utf-8")    # no expect()/assert - nothing to extract
    create_report(artifacts, tests, tmp_path / "report", combined=True)
    joined = chr(10).join(_text(tmp_path / "report" / "full-report.docx"))
    assert "(no assertion captured)" in joined


def test_no_test_summary_heading_when_there_is_nothing_to_summarise(tmp_path):
    from website_test_pipeline.report import RunReport, _test_summary_section
    from docx import Document as _Doc
    document = _Doc()
    _test_summary_section(document, RunReport())
    assert [p.text for p in document.paragraphs if p.style.name.startswith("Heading")] == []


def test_a_run_with_nothing_failing_says_so_plainly(tmp_path):
    artifacts, tests = _workspace(tmp_path)
    results = json.loads((artifacts / "test_results.json").read_text(encoding="utf-8"))
    results["tests"] = [r for r in results["tests"] if r["status"] == "passed"]
    (artifacts / "test_results.json").write_text(json.dumps(results), encoding="utf-8")
    (tmp_path / "flow_ratings.json").write_text(json.dumps({"version": 1, "ratings": {
        "f--good": [{"source": "runner", "at": "t1", "passed": True}]}}), encoding="utf-8")
    create_report(artifacts, tests, tmp_path / "report", combined=True)
    joined = chr(10).join(_text(tmp_path / "report" / "full-report.docx"))
    assert "No failures to triage this run." in joined
    assert "changed verdict across recent runs" not in joined


def test_evidence_caption_reads_like_a_person_wrote_it():
    from website_test_pipeline.report import _evidence_caption
    assert _evidence_caption("01-click-flow") == "Step 1: click flow"
    assert _evidence_caption("02_select_countrylist") == "Step 2: select countrylist"
    assert _evidence_caption("99-outcome") == "Final outcome"
    assert _evidence_caption("failure") == "failure"


def test_a_login_wall_is_reported_as_not_tested(tmp_path):
    artifacts, tests = _workspace(tmp_path)
    (artifacts / "login.inventory.json").write_text(json.dumps(
        {"url": "https://x.test/en/account", "auth": [{"kind": "login", "fields": [
            {"type": "email", "label": "Email"}, {"type": "password", "label": "Password"}]}]}), encoding="utf-8")
    create_report(artifacts, tests, tmp_path / "report", combined=True)
    joined = chr(10).join(_text(tmp_path / "report" / "full-report.docx"))
    assert "1 login or sign-up form(s) found" in joined and "Sign-in walls" in joined
    assert "https://x.test/en/account" in joined and "Email, Password" in joined


def test_no_login_wall_means_no_not_tested_block(tmp_path):
    artifacts, tests = _workspace(tmp_path)
    create_report(artifacts, tests, tmp_path / "report", combined=True)
    assert "Sign-in walls" not in chr(10).join(_text(tmp_path / "report" / "full-report.docx"))


def test_the_report_shows_the_authentication_line_and_a_finding_when_login_never_worked(tmp_path):
    artifacts, tests = _workspace(tmp_path)
    (artifacts / "auth_history.json").write_text(json.dumps([
        {"account": "default", "method": "popup", "status": "failed", "attempts": 3, "url": "https://x.test/login",
         "error": "Wrong password"},
    ]), encoding="utf-8")
    create_report(artifacts, tests, tmp_path / "report", combined=True)
    joined = chr(10).join(_text(tmp_path / "report" / "full-report.docx"))
    assert "Authentication:" in joined and "Could not sign in to the site" in joined
    assert "UNV‑001" in joined and "Open" in joined and "Test Owner" in joined and "Unassigned" not in joined


def test_the_report_shows_a_first_attempt_line_when_the_login_worked(tmp_path):
    artifacts, tests = _workspace(tmp_path)
    (artifacts / "auth_history.json").write_text(json.dumps([
        {"account": "default", "method": "env", "status": "signed-in-env", "attempts": 1, "url": "https://x.test/login"},
    ]), encoding="utf-8")
    create_report(artifacts, tests, tmp_path / "report", combined=True)
    joined = chr(10).join(_text(tmp_path / "report" / "full-report.docx"))
    assert "Signed in to the site on the first attempt (env)." in joined
    assert "login: default" not in joined                    # no auth_failure finding when it worked


def test_no_auth_history_means_no_authentication_line(tmp_path):
    artifacts, tests = _workspace(tmp_path)
    create_report(artifacts, tests, tmp_path / "report", combined=True)
    assert "Authentication:" not in chr(10).join(_text(tmp_path / "report" / "full-report.docx"))


# ------------------------------------------------------------------ plain-language translation (no raw Playwright
# code or pytest traces in the main-body tables; the Appendix keeps the raw form)

def test_plain_assertion_turns_the_known_expect_shapes_into_english():
    from website_test_pipeline.report import _plain_assertion
    cases = {
        "expect(page.get_by_role('heading', name='Login Page', exact=True)).to_be_visible()":
            'The "Login Page" heading should be visible on the page.',
        "expect(control).not_to_have_value('')": "It should not be left empty.",
        "expect(control).to_have_value('standard_user')": 'It should show "standard_user".',
        "expect(checkbox).to_be_checked()": "It should be checked.",
        "expect(button).to_be_disabled()": "It should be disabled.",
        "expect(page.locator('#inventory_container').first).to_have_count(1)":
            "1 of the element matching #inventory_container should be present.",
    }
    for code, plain in cases.items():
        assert _plain_assertion(code) == plain, code


def test_plain_assertion_preserves_the_real_case_of_a_name_unlike_str_capitalize():
    from website_test_pipeline.report import _plain_assertion
    out = _plain_assertion("expect(page.get_by_role('group', name='Passcode (default)', exact=True)).to_be_visible()")
    assert "Passcode (default)" in out and "passcode" not in out


def test_plain_assertion_reads_a_url_pattern_that_has_a_nested_closing_paren():
    from website_test_pipeline.report import _plain_assertion
    code = "expect(page).to_have_url(re.compile('/web/index\\.php/dashboard/index/?(?:[?#].*)?$'))"
    assert _plain_assertion(code) == "The page should end up at /web/index.php/dashboard/index."


def test_an_assertion_shape_not_in_the_known_vocabulary_falls_back_to_a_shortened_raw_line():
    from website_test_pipeline.report import _plain_assertion
    assert _plain_assertion("assert something_custom(x) == y") == "assert something_custom(x) == y"


def test_plain_error_covers_the_common_playwright_failure_shapes():
    from website_test_pipeline.findings import plain_error
    assert plain_error("E   AssertionError: Locator expected to be visible\nE   Error: element(s) not found") == \
        "The expected element never appeared on the page."
    assert "logout" not in plain_error(
        "E   AssertionError: Page URL expected to be 're.compile(\'/logout\')'\nE   Actual value: https://x.test/login")
    assert plain_error("E   playwright._impl._errors.TimeoutError: Locator.click: Timeout 30000ms exceeded.") == \
        "Timed out after 30s."
    assert "never found" in plain_error(
        'E   TimeoutError: Locator.click: Timeout 30000ms exceeded.\nE   Call log:\nE     - waiting for get_by_role("link", name="map")')
    assert "2 matching elements" in plain_error(
        'E   AssertionError: strict mode violation: locator("#x") resolved to 2 elements')
    assert plain_error(None) == "failed with no captured reason"


def test_the_test_execution_summary_and_defect_report_contain_no_raw_playwright_code(tmp_path):
    artifacts, tests = _workspace(tmp_path)
    create_report(artifacts, tests, tmp_path / "report", combined=True)
    lines = _text(tmp_path / "report" / "full-report.docx")
    main_body = chr(10).join(lines[:_last_index(lines, "Test Logs & Evidence")])   # the evidence section keeps the raw form on purpose
    for jargon in ("get_by_role(", "expect(", "Locator.", "playwright._impl", "AssertionError:"):
        assert jargon not in main_body, jargon


# ------------------------------------------------------------------ "Needs a human" (CAPTCHA / verification code)

def test_a_captcha_page_is_listed_as_needing_a_human(tmp_path):
    artifacts, tests = _workspace(tmp_path)
    (artifacts / "subscribe.inventory.json").write_text(json.dumps(
        {"url": "https://x.test/subscribe", "controls": [
            {"name": "What code is in the image?", "tag": "input", "selector": "#edit-captcha-response",
             "field_name": "captcha_response"}],
         "forms": [], "accessibility": ""}), encoding="utf-8")
    create_report(artifacts, tests, tmp_path / "report", combined=True)
    joined = chr(10).join(_text(tmp_path / "report" / "full-report.docx"))
    assert "Needs a human" in joined and "CAPTCHA" in joined and "https://x.test/subscribe" in joined


def test_a_verification_code_field_is_also_flagged(tmp_path):
    from website_test_pipeline.findings import human_input_pages
    inv = [{"url": "https://x.test/2fa", "controls": [{"name": "Enter your verification code", "tag": "input"}],
           "forms": [], "accessibility": ""}]
    pages = human_input_pages(inv)
    assert len(pages) == 1 and pages[0]["reason"] == "verification code"


def test_an_ordinary_page_is_never_flagged_as_needing_a_human(tmp_path):
    artifacts, tests = _workspace(tmp_path)
    create_report(artifacts, tests, tmp_path / "report", combined=True)
    assert "Needs a human" not in chr(10).join(_text(tmp_path / "report" / "full-report.docx"))


# ------------------------------------------------------------------ a page test the validator deliberately
# skipped (not a failure) - must not be counted or shown as one

def test_a_skipped_page_test_is_its_own_verdict_not_a_failure(tmp_path):
    artifacts, tests = _workspace(tmp_path)
    skip_spec = tests / "https-x-test-skip_test.py"
    skip_spec.write_text("def test_page_not_testable(page): pass\n", encoding="utf-8")
    skip_node = "tests/https-x-test-skip_test.py::test_page_not_testable[chromium]"
    results = json.loads((artifacts / "test_results.json").read_text(encoding="utf-8"))
    results["tests"].append({
        "nodeid": skip_node, "title": "", "url": START + "skip", "status": "skipped", "duration": 0.0,
        "error": ("('C:\\repo\\tests\\https-x-test-skip_test.py', 5, \"Skipped: NOT TESTABLE: no spec passed "
                  "validation - unstable text selector 'get_by_text'\")")})
    (artifacts / "test_results.json").write_text(json.dumps(results), encoding="utf-8")
    create_report(artifacts, tests, tmp_path / "report", flows_file=tmp_path / "flows.json",
                 ratings_file=tmp_path / "flow_ratings.json", combined=True)
    joined = chr(10).join(_text(tmp_path / "report" / "full-report.docx"))
    assert "SKIPPED" in joined
    assert "Not tested: NOT TESTABLE" in joined
    # the raw skip tuple never appears anywhere in the main body
    lines = _text(tmp_path / "report" / "full-report.docx")
    main_body = chr(10).join(lines[:_last_index(lines, "Test Logs & Evidence")])
    assert "C:\\repo" not in main_body and "get_by_text" not in main_body


def test_a_skipped_test_is_excluded_from_the_pass_rate_and_failed_count(tmp_path):
    artifacts, tests = _workspace(tmp_path)
    skip_spec = tests / "https-x-test-skip_test.py"
    skip_spec.write_text("def test_page_not_testable(page): pass\n", encoding="utf-8")
    results = json.loads((artifacts / "test_results.json").read_text(encoding="utf-8"))
    results["tests"].append({
        "nodeid": "tests/https-x-test-skip_test.py::test_page_not_testable[chromium]", "title": "",
        "url": START + "skip", "status": "skipped", "duration": 0.0,
        "error": "('f.py', 1, \"Skipped: no reason\")"})
    (artifacts / "test_results.json").write_text(json.dumps(results), encoding="utf-8")
    create_report(artifacts, tests, tmp_path / "report", flows_file=tmp_path / "flows.json",
                 ratings_file=tmp_path / "flow_ratings.json", combined=True)
    joined = chr(10).join(_text(tmp_path / "report" / "full-report.docx"))
    assert "2 tests ran (0 page + 2 flow): 2 passed and 0 failed. 1 were skipped by the harness and 1 were invalid" in joined


# ------------------------------------------------------------------ Defect Report: steps to reproduce, and a
# working link from the bug's Title straight to its own evidence in the Appendix

def test_the_defect_report_has_steps_to_reproduce_and_a_reason_when_evidence_is_missing(tmp_path):
    artifacts, tests = _workspace(tmp_path)
    create_report(artifacts, tests, tmp_path / "report", combined=True)
    doc = Document(str(tmp_path / "report" / "full-report.docx"))
    table = next(t for t in doc.tables if [c.text for c in t.rows[0].cells] == ["ID", "Kind", "Title", "URL", "What happened", "Evidence"])
    row = table.rows[1]
    assert row.cells[0].text == "TST‑001" and row.cells[1].text == "Test defect"
    assert row.cells[2].text == "wizard step"                       # the full title, never cut
    detail = next(t for t in doc.tables if [c.text for c in t.rows[0].cells] == ["Field", "Details"])
    fields = {r.cells[0].text: r.cells[1].text for r in detail.rows[1:]}
    assert "Open https://x.test/en." in fields["Steps to reproduce"]
    assert "Run test tests/https-x-test-en_test.py::test_wizard_step[chromium]." in fields["Steps to reproduce"]
    assert fields["Severity (how bad)"] == "Low" and fields["Owner"] == "Test Owner"
    assert fields["Evidence"].startswith("No screenshot")           # no link, so the literal reason
    assert len(fields["Due"]) == 10 and fields["Due"][4] == "-"      # an ISO date


def test_the_report_is_refused_when_nobody_is_named_as_owner(tmp_path, monkeypatch):
    import pytest
    from website_test_pipeline.report_policy import ReportConsistencyError
    monkeypatch.delenv("REPORT_OWNER", raising=False)
    artifacts, tests = _workspace(tmp_path)
    with pytest.raises(ReportConsistencyError, match="no owner"):
        create_report(artifacts, tests, tmp_path / "report", combined=True)
    (tmp_path / "owners.json").write_text(json.dumps({"default": "Dana Reviewer"}), encoding="utf-8")
    create_report(artifacts, tests, tmp_path / "report", combined=True)


def test_a_non_blocking_open_defect_states_why_it_does_not_block():
    """Part 8: "No release blockers were recorded" next to a report with one open Medium defect reads as a
    contradiction unless it says why that defect isn't a blocker."""
    from website_test_pipeline.report import _risks_section
    report_data = {"defects": [{"id": "DEF-001", "category": "defect", "status": "Open", "severity": "Medium",
                                "title": "t", "blocking": False, "business_impact": "x"}],
                   "taxonomy": {"defect": 1, "unverified": 0, "limitation": 0, "tooling": 0, "test_defect": 0},
                   "risks": []}
    document = Document()
    _risks_section(document, None, report_data)
    text = "\n".join(p.text for p in document.paragraphs)
    assert "No release blockers were recorded. DEF-001 is Medium, below the High/Critical blocking threshold" in text


def test_find_attachments_copies_into_the_runs_own_folder_instead_of_linking_the_temp_source(tmp_path):
    """Part 4: a trace/video link into pytest-playwright's --output temp dir resolves for nobody and is wiped
    on reboot - the fix is to copy the artifact into the run's own evidence folder (the same place and
    pattern Screenshot links already use) and return THAT path, not the original temp one."""
    pw_dir = tmp_path / "pw_temp" / "some-test-chromium"
    pw_dir.mkdir(parents=True)
    source = pw_dir / "trace.zip"
    source.write_bytes(b"fake trace data")
    dest_dir = tmp_path / "run_evidence" / "some-test"
    found = _find_attachments(tmp_path / "pw_temp", "some-test", copy_into=dest_dir)
    assert len(found) == 1
    assert Path(found[0]).parent == dest_dir   # points into the run's own tree, not pw_temp
    assert Path(found[0]).read_bytes() == b"fake trace data"
    assert "pw_temp" not in found[0]


def test_a_defect_that_passes_this_run_still_renders_as_open_until_a_person_closes_it(tmp_path):
    """Part 7: the ledger mechanism that is supposed to keep a defect visible after a passing run has to be
    exercised against an ACTUAL render, not just the ledger functions in isolation - f--good passes clean
    this run (test_results.json below), but a previous run's failure is already in defect_ledger.json."""
    from website_test_pipeline import report_policy as policy
    artifacts, tests = _workspace(tmp_path)
    good_node = f"tests/{file_name(_flow('f--good'))}::test_flow_good[chromium]"
    key = policy.stable_test_key(good_node)
    ledger = {"version": 1, "defects": {key: {
        "title": "A visitor picks a country and sees the results.", "url": START, "scope": "flow",
        "first_seen_run": "run-1", "first_seen_at": "2026-10-01T00:00:00+00:00",
        "last_failed_run": "run-1", "last_failed_at": "2026-10-01T00:00:00+00:00",
        "total_failures": 3, "consecutive_not_reproduced": 0, "status": "open",
        "closed_by": None, "closed_at": None, "closed_reason": None, "closure_history": []}}}
    (tmp_path / "defect_ledger.json").write_text(json.dumps(ledger), encoding="utf-8")
    create_report(artifacts, tests, tmp_path / "report", combined=True)
    joined = chr(10).join(_text(tmp_path / "report" / "full-report.docx"))
    # C3: a not-reproduced ghost was not found THIS run - it is carried from a previous one.
    assert "1 open defect carried from previous runs (not reproduced this run)" in joined
    assert "0 application defects found this run" in joined
    assert "not reproduced this run, confirmed in 3 prior report(s), still open" in joined
    assert "DEF" in joined  # has a row in the Defects table under a DEF- id, not dropped from it
    updated = json.loads((tmp_path / "defect_ledger.json").read_text(encoding="utf-8"))
    assert updated["defects"][key]["consecutive_not_reproduced"] == 1
    assert updated["defects"][key]["status"] == "open"   # still open - nothing but an explicit close removes it


def test_an_aborted_build_never_commits_a_ledger_mutation(tmp_path, monkeypatch):
    """P1 root cause: report.py used to call save_ledger (inside _defect_records) BEFORE validate_report_data
    could reject the build - a build that aborted after that point still permanently incremented
    consecutive_not_reproduced, which is exactly how a real defect's auto-close fired one run early. The
    ledger write must be deferred until the whole build is proven to succeed, the same way save_history is."""
    import website_test_pipeline.report as report_mod
    from website_test_pipeline import report_policy as policy
    artifacts, tests = _workspace(tmp_path)
    good_node = f"tests/{file_name(_flow('f--good'))}::test_flow_good[chromium]"
    key = policy.stable_test_key(good_node)
    ledger = {"version": 1, "defects": {key: {
        "title": "t", "url": START, "scope": "flow", "first_seen_run": "run-1", "first_seen_at": "2026-10-01T00:00:00+00:00",
        "last_failed_run": "run-1", "last_failed_at": "2026-10-01T00:00:00+00:00",
        "total_failures": 1, "consecutive_not_reproduced": 1, "status": "open",
        "closed_by": None, "closed_at": None, "closed_reason": None, "closure_history": []}}}
    ledger_path = tmp_path / "defect_ledger.json"
    ledger_path.write_text(json.dumps(ledger), encoding="utf-8")
    before = ledger_path.read_text(encoding="utf-8")

    def _boom(data):
        raise ValueError("forced validation failure for this test")
    monkeypatch.setattr(report_mod, "validate_report_data", _boom)
    with pytest.raises(ValueError, match="forced validation failure"):
        create_report(artifacts, tests, tmp_path / "report", combined=True)
    assert ledger_path.read_text(encoding="utf-8") == before   # byte-for-byte unchanged - no mutation committed


def test_a_defect_at_two_consecutive_clean_runs_renders_open_with_the_counter_visible(tmp_path):
    """P1.5: proof the suite would have caught the premature close - a defect that reaches 2 of 3 consecutive
    clean runs must still render as OPEN with the progress counter visible; only the 3rd closes it, never
    earlier (confirmed live: a build that aborted after writing the ledger let a 2-of-3 state masquerade as
    the 3rd clean run, closing the real DEF-001 one run early)."""
    from website_test_pipeline import report_policy as policy
    artifacts, tests = _workspace(tmp_path)
    good_node = f"tests/{file_name(_flow('f--good'))}::test_flow_good[chromium]"
    key = policy.stable_test_key(good_node)
    ledger = {"version": 1, "defects": {key: {
        "title": "A visitor picks a country and sees the results.", "url": START, "scope": "flow",
        "first_seen_run": "run-1", "first_seen_at": "2026-10-01T00:00:00+00:00",
        "last_failed_run": "run-1", "last_failed_at": "2026-10-01T00:00:00+00:00",
        "total_failures": 3, "consecutive_not_reproduced": 1, "status": "open",
        "closed_by": None, "closed_at": None, "closed_reason": None, "closure_history": []}}}
    (tmp_path / "defect_ledger.json").write_text(json.dumps(ledger), encoding="utf-8")
    create_report(artifacts, tests, tmp_path / "report", combined=True)   # this run: 1 -> 2
    joined = chr(10).join(_text(tmp_path / "report" / "full-report.docx"))
    assert "consecutive clean runs: 2 of 3" in joined
    assert "still open" in joined
    assert "DEF" in joined   # still has a row in the Defects table - not closed, not absent
    after_2 = json.loads((tmp_path / "defect_ledger.json").read_text(encoding="utf-8"))
    assert after_2["defects"][key]["status"] == "open" and after_2["defects"][key]["consecutive_not_reproduced"] == 2
    create_report(artifacts, tests, tmp_path / "report", combined=True)   # this run: 2 -> 3, closes
    after_3 = json.loads((tmp_path / "defect_ledger.json").read_text(encoding="utf-8"))
    assert after_3["defects"][key]["status"] == "closed"
    assert after_3["defects"][key]["audit_log"][-1]["actor"] == "[auto]"
    assert after_3["defects"][key]["audit_log"][-1]["clean_run_count"] == 3
    final_data = json.loads((tmp_path / "report" / "report-data.json").read_text(encoding="utf-8"))
    assert final_data["counts"]["open_defects"] == 0
    assert not any(d["status"] == "Open" and d["category"] == "defect" for d in final_data["defects"])


def test_a_second_run_gets_a_delta_block_with_removed_tests_and_their_reason(tmp_path):
    artifacts, tests = _workspace(tmp_path)
    create_report(artifacts, tests, tmp_path / "report", combined=True)
    first = chr(10).join(_text(tmp_path / "report" / "full-report.docx"))
    assert "this run is the baseline" in first
    results = json.loads((artifacts / "test_results.json").read_text(encoding="utf-8"))
    results["finished_at"] = "later"
    results["tests"] = [r for r in results["tests"] if "test_flow_flap" not in r["nodeid"]]
    (artifacts / "test_results.json").write_text(json.dumps(results), encoding="utf-8")
    create_report(artifacts, tests, tmp_path / "report", combined=True)
    second = chr(10).join(_text(tmp_path / "report" / "full-report.docx"))
    assert "Change since the previous run" in second and "Compared with" in second
    assert "Tests removed" in second and "test_flow_flap" in second


def test_a_defect_lives_its_whole_life_through_real_renders_open_persist_frozen_severity_then_auto_close(tmp_path):
    """The ledger tests above start from a hand-written ledger entry. This one starts from nothing: a flow FAILS
    in a real create_report, the ledger entry is born, it survives a second failure with its severity frozen,
    and only the third consecutive clean run closes it - each step through the actual build, not the helpers."""
    from website_test_pipeline import report_policy as policy
    artifacts, tests = _workspace(tmp_path)
    good_node = f"tests/{file_name(_flow('f--good'))}::test_flow_good[chromium]"
    key = policy.stable_test_key(good_node)
    ledger_path = tmp_path / "defect_ledger.json"

    def run(status, error=None):
        results = json.loads((artifacts / "test_results.json").read_text(encoding="utf-8"))
        for row in results["tests"]:
            if row["nodeid"] == good_node:
                row.update(status=status, error=error)
        (artifacts / "test_results.json").write_text(json.dumps(results), encoding="utf-8")
        create_report(artifacts, tests, tmp_path / "report", combined=True)
        entry = json.loads(ledger_path.read_text(encoding="utf-8"))["defects"].get(key)
        return entry, json.loads((tmp_path / "report" / "report-data.json").read_text(encoding="utf-8"))

    wrong_page = "E   AssertionError: Page.expect_url\nE     Expected URL to match /en/find but was https://x.test/en"
    entry, data = run("failed", wrong_page)                              # 1: born from a real failure
    assert entry["status"] == "open" and entry["total_failures"] == 1 and entry["severity"] == "High"
    assert data["recommendation"] == "NO-GO" and data["counts"]["open_defects"] == 1
    born_severity, born_at = entry["severity"], entry["severity_set_at"]

    entry, data = run("failed", wrong_page.replace("/en/find", "/en/other"))   # 2: fails again, differently
    assert entry["total_failures"] == 2 and entry["consecutive_not_reproduced"] == 0
    assert (entry["severity"], entry["severity_set_at"]) == (born_severity, born_at)   # frozen: set once, at confirmation

    for clean in (1, 2):                                                 # 3-4: clean runs do not close it
        entry, data = run("passed")
        assert entry["status"] == "open" and entry["consecutive_not_reproduced"] == clean
        assert data["counts"]["open_defects"] == 1 and data["recommendation"] == "NO-GO"

    entry, data = run("passed")                                          # 5: the third consecutive clean run closes it
    assert entry["status"] == "closed" and entry["audit_log"][-1]["actor"] == "[auto]"
    assert entry["audit_log"][-1]["clean_run_count"] == 3 and entry["severity"] == born_severity
    assert data["counts"]["open_defects"] == 0 and data["recommendation"] != "NO-GO"


def test_two_real_defects_of_equal_weight_still_produce_a_report_that_says_so(tmp_path):
    """Found live: a run with three genuine, equally weighted findings was refused outright ("every scored item has the same
    priority"), so it produced no report at all. A tie is now flagged in "Validation warnings" for a person to confirm."""
    artifacts, tests = _workspace(tmp_path)
    results = json.loads((artifacts / "test_results.json").read_text(encoding="utf-8"))
    wrong_page = "E   AssertionError: Page.expect_url\nE     Expected URL to match /en/find but was https://x.test/en"
    for row in results["tests"]:
        if "test_flow_" in row["nodeid"]:
            row.update(status="failed", error=wrong_page)
    (artifacts / "test_results.json").write_text(json.dumps(results), encoding="utf-8")
    create_report(artifacts, tests, tmp_path / "report", combined=True)                  # used to raise ReportConsistencyError
    data = json.loads((tmp_path / "report" / "report-data.json").read_text(encoding="utf-8"))
    scored = [d for d in data["defects"] if d["category"] in {"defect", "unverified"}]
    assert len(scored) == 2 and len({d["priority"] for d in scored}) == 1
    text = chr(10).join(_text(tmp_path / "report" / "full-report.docx"))
    assert "Scoring check" in text and "All 2 scored items" in text and "judged from its own user impact" in text


def test_a_check_that_could_not_click_is_triage_not_a_defect_and_the_verdict_stays_go(tmp_path):
    """Found live on a production site: a hidden skip link (covered by the header) was reported as a High defect -> NO-GO."""
    artifacts, tests = _workspace(tmp_path)
    results = json.loads((artifacts / "test_results.json").read_text(encoding="utf-8"))
    for row in results["tests"]:
        if "test_flow_good" in row["nodeid"]:
            row.update(status="failed", error="E   playwright._impl._errors.TimeoutError: Locator.click: Timeout 30000ms exceeded.\n"
                                              "E   Call log:\nE     - <div class=\"header\"></div> intercepts pointer events")
    (artifacts / "test_results.json").write_text(json.dumps(results), encoding="utf-8")
    create_report(artifacts, tests, tmp_path / "report", combined=True)
    joined = chr(10).join(_text(tmp_path / "report" / "full-report.docx"))
    assert "0 application defects found this run." in joined
    assert "TRI‑001" in joined and "Needs triage" in joined
    assert "Release decision: GO" in joined
    data = json.loads((tmp_path / "report" / "report-data.json").read_text(encoding="utf-8"))
    assert data["recommendation"] == "GO" and data["taxonomy"]["triage"] == 1 and data["counts"]["failed"] == 1


def test_a_run_where_no_test_ran_has_no_verdict(tmp_path):
    """Found live: every model call was refused, nothing was generated, and the report still said CONDITIONAL GO."""
    artifacts, tests = _workspace(tmp_path)
    (artifacts / "test_results.json").write_text(json.dumps({"finished_at": "t", "tests": []}), encoding="utf-8")
    create_report(artifacts, tests, tmp_path / "report", combined=True)
    joined = chr(10).join(_text(tmp_path / "report" / "full-report.docx"))
    data = json.loads((tmp_path / "report" / "report-data.json").read_text(encoding="utf-8"))
    assert data["recommendation"] == "NO VERDICT"
    assert "Final verdict: NO VERDICT: no test ran, so this run has no evidence either way." in joined
    assert "NOT ASSESSED" in joined and "CONDITIONAL GO" not in joined.split("Release rule:")[0]
