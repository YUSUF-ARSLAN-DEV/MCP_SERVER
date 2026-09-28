import pytest
from docx import Document

from website_test_pipeline.report import _validate_docx_export
from website_test_pipeline.report_model import canonical_counts, validate_report_data


class _Outcome:
    def __init__(self, status):
        self.status = status


class _Report:
    def __init__(self, outcomes):
        self.outcomes = outcomes


class _Run:
    url_reports = [_Report([_Outcome("passed"), _Outcome("failed"), _Outcome("skipped")])]
    tested_flows = []
    untested_flows = [object()]


def test_canonical_counts_reconcile_and_exclude_skips_from_pass_rate():
    counts = canonical_counts(_Run())
    assert counts == {
        "executed_tests": 2,
        "passed": 1,
        "failed": 1,
        "skipped": 1,
        "blocked_flows": 1,
        "all_items": 4,
        "pass_rate_percent": 50,
    }


def test_conditional_go_requires_a_linked_condition():
    data = {
        "schema_version": "1.0", "generated_at": "2026-01-01T00:00:00+00:00",
        "counts": {"executed_tests": 0, "passed": 0, "failed": 0, "skipped": 0,
                   "blocked_flows": 0, "all_items": 0, "pass_rate_percent": 0},
        "recommendation": "CONDITIONAL GO", "conditions": [], "defects": [], "coverage": {},
    }
    with pytest.raises(ValueError, match="CONDITIONAL GO"):
        validate_report_data(data)


def test_docx_export_rejects_the_known_repeated_value_corruption(tmp_path):
    document = Document()
    document.add_paragraph(" ".join(["1"] * 40))
    path = tmp_path / "corrupt.docx"
    document.save(path)
    with pytest.raises(ValueError, match="repeated-value corruption"):
        _validate_docx_export(path)


def test_docx_export_accepts_normal_prose(tmp_path):
    document = Document()
    document.add_paragraph("The test run completed with a clear result.")
    path = tmp_path / "valid.docx"
    document.save(path)
    _validate_docx_export(path)
