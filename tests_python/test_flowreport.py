import json
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from PIL import Image

from website_test_pipeline.flowgen import file_name
from website_test_pipeline.flowreport import (
    attribute_failure, build_flow_report, completed_steps, flow_warnings, pages_touched,
)
from website_test_pipeline.report import RunReport, _save_document, _slug, create_report, load_run, name_for

START = "https://x.test/"
GOAL = "A visitor picks a country and searches, then sees the frequencies page."


def _flow(fid="x--search", status="verified", **over):
    flow = {
        "id": fid, "goal": GOAL, "source": "intent", "status": status, "start_url": START, "intent_id": "i-001",
        "steps": [{"kind": "select", "selector": "#country", "name": "Country", "value": "Egypt"},
                  {"kind": "click", "selector": None, "name": "Search"},
                  {"kind": "click", "selector": None, "name": "Details"}],
        "outcome": {"effect": "navigates", "to": "/en/find"},
        "observed": {"effect": "navigates", "url": "https://x.test/en/find", "landed_url": "https://x.test/en",
                     "new_headings": ["Find frequencies"], "step_effects": ["no-visible-change", "navigates", "reveals"],
                     "step_urls": ["https://x.test/en", "https://x.test/en/find", "https://x.test/en/find"]},
    }
    flow.update(over)
    return flow


def _outcome(status="passed", evidence=(), error=None, attachments=()):
    return SimpleNamespace(status=status, evidence=list(evidence), error=error, attachments=list(attachments),
                           passed=status == "passed", duration=1.5, assertions=[], title="t")


# ------------------------------------------------------------------ pure logic

def test_completed_steps_counts_the_unbroken_run_of_step_screenshots():
    assert completed_steps(["01-select-country", "02-click-search", "99-outcome"]) == 2
    assert completed_steps(["01-a", "03-c"]) == 1            # step 2 has no screenshot: it is the one that broke
    assert completed_steps([]) == 0
    assert completed_steps(["99-outcome"]) == 0


def test_pages_touched_follows_the_journey_without_repeats():
    assert pages_touched(_flow()) == ["/en", "/en/find"]
    assert pages_touched({"start_url": START}) == ["/"]              # an old run with no recorded steps


def test_the_failure_names_the_step_the_page_it_should_reach_and_where_the_browser_was():
    error = "stuff\nE   AssertionError: Page URL expected to match\nE   Actual value: https://x.test/en/other\n"
    text = attribute_failure(_flow(), ["01-select-country"], error)
    assert text.startswith('Step 2 of 3 did not complete: click "Search".')
    assert "should have ended on /en/find" in text and "browser was on /en/other" in text
    assert "Reason: AssertionError: Page URL expected to match" in text


def test_a_failure_after_every_step_is_an_outcome_check_failure():
    text = attribute_failure(_flow(), ["01-a", "02-b", "03-c"], "E   heading not visible")
    assert text.startswith("All 3 steps ran, but the final outcome check failed.") and "heading not visible" in text


def test_a_setup_error_is_not_blamed_on_a_step():
    assert "did not start" in attribute_failure(_flow(), [], "boom", status="error")


def test_flow_report_carries_the_journey_and_the_history():
    entries = [{"source": "runner", "at": "2026-09-20T09:00:00", "passed": True,
                "checks": {"steps_completed": "3/3", "observed_effect": "navigates"}},
               {"source": "human", "at": "2026-09-20T10:00:00", "by": "sam", "decision": "approved", "reason": "ok"}]
    fr = build_flow_report(_flow(), entries, _outcome())
    assert fr.title == GOAL and fr.intent_id == "i-001" and fr.tested and fr.passed
    assert fr.steps[0] == 'select "Egypt" in #country' and fr.lands == ["/en", "/en/find", "/en/find"]
    assert fr.pages == ["/en", "/en/find"] and fr.expected == "navigates -> /en/find" and fr.observed == "navigates at /en/find"
    assert fr.new_headings == ["Find frequencies"] and "[human sam] approved - ok" in fr.history[1]
    assert fr.failure == ""


def test_landing_pages_are_dropped_when_they_do_not_match_the_steps():
    flow = _flow()
    flow["observed"]["step_urls"] = ["https://x.test/en"]
    assert build_flow_report(flow, [], _outcome()).lands == []


def test_a_flow_that_did_not_run_has_no_outcome_and_no_warnings():
    fr = build_flow_report(_flow(), [], None)
    assert not fr.tested and not fr.passed and not fr.failed and fr.warnings == []


def test_warnings_flag_missing_evidence_a_tolerated_failure_and_staleness():
    assert any("no screenshot evidence" in w for w in flow_warnings(build_flow_report(_flow(), [], _outcome())))
    failed = build_flow_report(_flow(), [], _outcome("failed", evidence=["01-a.png"]))
    assert any("still verified" in w for w in failed.warnings)
    assert any("stale" in w for w in build_flow_report(_flow(status="stale"), [], _outcome("failed", evidence=["01-a.png"])).warnings)
    assert flow_warnings(build_flow_report(_flow("x--y", "approved"), [], _outcome("passed", evidence=["01-a.png"]))) == []


# ------------------------------------------------------------------ the documents

def _png(path: Path) -> str:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (20, 20), "white").save(path)
    return str(path)


def _workspace(tmp_path, with_flows=True, flow_status="failed"):
    artifacts, tests = tmp_path / "artifacts", tmp_path / "tests"
    artifacts.mkdir()
    tests.mkdir()
    flow = _flow()
    flow_node = f"tests/{file_name(flow)}::test_flow_x_search[chromium]"
    page_node = "tests/https-x-test_test.py::test_home_heading[chromium]"
    (tests / file_name(flow)).write_text(
        "from playwright.sync_api import expect" + chr(10) + "def test_flow_x_search(page):" + chr(10)
        + "    expect(page).to_have_url('x')" + chr(10), encoding="utf-8")
    rows = [{"nodeid": page_node, "title": "home heading", "url": START, "status": "passed", "duration": 0.5, "error": None},
            {"nodeid": flow_node, "title": "", "url": START, "status": flow_status, "duration": 2.0,
             "error": None if flow_status == "passed" else "E   AssertionError: Page URL expected" + chr(10) + "E   Actual value: https://x.test/en/other"}]
    (artifacts / "test_results.json").write_text(json.dumps({"finished_at": "t", "tests": rows}), encoding="utf-8")
    _png(artifacts / "evidence" / _slug(page_node) / "01-home.png")
    _png(artifacts / "evidence" / _slug(flow_node) / "01-select-country.png")
    if flow_status == "passed":
        _png(artifacts / "evidence" / _slug(flow_node) / "02-click-search.png")
    if with_flows:
        untested = _flow("x--map", "candidate", goal="A visitor opens the map.")
        (tmp_path / "flows.json").write_text(json.dumps({"version": 1, "flows": [flow, untested]}), encoding="utf-8")
        (tmp_path / "flow_ratings.json").write_text(json.dumps({"version": 1, "ratings": {"x--search": [
            {"source": "human", "at": "2026-09-20T10:00:00", "by": "sam", "decision": "approved", "reason": "core"}]}}), encoding="utf-8")
    return artifacts, tests


def _text(path: Path) -> list[str]:
    doc = Document(str(path))
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        lines += [cell.text for row in table.rows for cell in row.cells]
    return lines


def test_the_flow_test_leaves_the_page_tests_and_is_counted_once(tmp_path):
    artifacts, tests = _workspace(tmp_path)
    run = load_run(artifacts, tests)
    assert [o.title for u in run.url_reports for o in u.outcomes] == ["home heading"]
    assert [f.flow_id for f in run.tested_flows] == ["x--search"] and [f.flow_id for f in run.untested_flows] == ["x--map"]
    assert (run.total, run.passed, run.failed) == (2, 1, 1)
    assert run.flows_at(START)[0].failure.startswith("Step 2 of 3 did not complete")


def test_the_combined_report_has_a_user_flows_section_in_plain_language(tmp_path):
    artifacts, tests = _workspace(tmp_path)
    create_report(artifacts, tests, tmp_path / "report", combined=True)
    text = _text(tmp_path / "report" / "full-report.docx")
    joined = chr(10).join(text)
    for part in (GOAL, "User flows", "Journey", "Expected vs observed", "Where it broke", "Review history",
                 "Step 2 of 3 did not complete", "[human sam] approved - core", "Flows without a test result in this run",
                 "A visitor opens the map.", "/en → /en/find"):
        assert part in joined, part
    assert 'select "Egypt" in #country' in joined


def test_the_page_document_of_the_start_url_carries_its_flows(tmp_path):
    artifacts, tests = _workspace(tmp_path)
    create_report(artifacts, tests, tmp_path / "report")
    joined = chr(10).join(_text(tmp_path / "report" / (name_for(START) + ".docx")))
    assert GOAL in joined and "home heading" in joined.replace("_", " ")
    assert "still verified" in joined                 # the tolerated-failure warning reaches the page document


def test_a_passing_flow_shows_no_failure_section(tmp_path):
    artifacts, tests = _workspace(tmp_path, flow_status="passed")
    create_report(artifacts, tests, tmp_path / "report", combined=True)
    joined = chr(10).join(_text(tmp_path / "report" / "full-report.docx"))
    assert GOAL in joined and "Where it broke" not in joined and "Result: PASSED" in joined


def test_without_a_flows_file_the_flow_test_stays_an_ordinary_test(tmp_path):
    artifacts, tests = _workspace(tmp_path, with_flows=False)
    run = load_run(artifacts, tests)
    assert run.flow_reports == [] and sum(u.total for u in run.url_reports) == 2 and run.total == 2
    create_report(artifacts, tests, tmp_path / "report", combined=True)
    assert "User flows" not in chr(10).join(_text(tmp_path / "report" / "full-report.docx"))


def test_an_unreadable_flows_file_never_breaks_the_report(tmp_path):
    artifacts, tests = _workspace(tmp_path, with_flows=False)
    (tmp_path / "flows.json").write_text("not json", encoding="utf-8")
    assert load_run(artifacts, tests).total == 2


def test_a_flow_that_starts_on_a_page_with_no_page_tests_still_gets_a_document(tmp_path):
    artifacts, tests = _workspace(tmp_path)
    rows = json.loads((artifacts / "test_results.json").read_text(encoding="utf-8"))
    rows["tests"] = rows["tests"][1:]                            # drop the page test: only the flow ran
    (artifacts / "test_results.json").write_text(json.dumps(rows), encoding="utf-8")
    create_report(artifacts, tests, tmp_path / "report")
    assert (tmp_path / "report" / (name_for(START) + ".docx")).exists()


def test_a_locked_document_is_saved_under_a_new_name_instead_of_aborting_the_report(tmp_path):
    class _Doc:
        def __init__(self):
            self.saved = []

        def save(self, path):
            if Path(path).name == "locked.docx":
                raise PermissionError("open in Word")
            self.saved.append(Path(path).name)

    run, doc = RunReport(), _Doc()
    assert _save_document(doc, tmp_path / "free.docx", run).name == "free.docx"
    assert _save_document(doc, tmp_path / "locked.docx", run).name == "locked-new.docx"
    assert doc.saved == ["free.docx", "locked-new.docx"]
    assert len(run.notes) == 1 and "locked.docx" in run.notes[0] and run.notes[0] in run.warnings


def test_a_passing_flow_that_only_proves_a_url_change_is_flagged_as_weak():
    flow = _flow()
    flow["observed"].update(new_headings=[], results=[], new_controls=[])
    weak = build_flow_report(flow, [], _outcome("passed", evidence=["01-a.png"]))
    assert weak.navigation_only and any("only proves the URL changed" in w for w in weak.warnings)
    strong = build_flow_report(_flow(), [], _outcome("passed", evidence=["01-a.png"]))
    assert not strong.navigation_only and not any("only proves" in w for w in strong.warnings)
    reveals = _flow()
    reveals["observed"].update(effect="reveals", new_headings=[])
    assert not build_flow_report(reveals, [], _outcome("passed", evidence=["01-a.png"])).navigation_only
